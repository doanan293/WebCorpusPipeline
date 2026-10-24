from __future__ import annotations

from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from data_pipeline.domain.reader.models import DownloadedFile, ExtractedPart, LocalResult
from data_pipeline.infrastructure.reader.readers.common import result


def _shape_text(shape: Any) -> list[str]:
    out = []
    if getattr(shape, "has_text_frame", False) and shape.text.strip():
        out.append(shape.text.strip())
    if getattr(shape, "has_table", False):
        out.append(
            "\n".join(
                "\t".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows
            )
        )
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            out.extend(_shape_text(child))
    return out


def _has_picture(shapes: Any) -> bool:
    return any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        or (shape.shape_type == MSO_SHAPE_TYPE.GROUP and _has_picture(shape.shapes))
        for shape in shapes
    )


def parse_powerpoint(file: DownloadedFile) -> LocalResult:
    deck = Presentation(str(file.path))
    parts = []
    for index, slide in enumerate(deck.slides, 1):
        chunks = [text for shape in slide.shapes for text in _shape_text(shape)]
        parts.append(
            ExtractedPart(
                index, "slide", "\n".join(chunks), _has_picture(slide.shapes), "python-pptx"
            )
        )
    return result(file, parts, "python-pptx")
