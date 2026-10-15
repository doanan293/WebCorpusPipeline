import base64
import hashlib
from pathlib import Path

import pymupdf as fitz
import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from data_pipeline.enrich.download.artifacts import extension_for
from data_pipeline.enrich.models import DownloadedFile
from data_pipeline.enrich.parsing.local import dispatch, parse_local
from data_pipeline.enrich.parsing.local.pdf import cloud_part_indices, parse_pdf


def _file(path: Path) -> DownloadedFile:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    return DownloadedFile(
        path.as_uri(), path, digest, path.stat().st_size, "application/octet-stream", path.as_uri()
    )


def test_pdf_keeps_short_text_and_only_visual_pages_are_cloud_candidates(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "A 0")
    doc.new_page()
    doc.save(str(path))
    doc.close()
    parsed = parse_pdf(_file(path))
    assert parsed.parts[0].text == "A 0"
    assert cloud_part_indices(parsed) == []


def test_word_reads_paragraphs_and_tables_in_plain_text(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Heading")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    doc.add_paragraph("Body")
    doc.save(str(path))
    result = parse_local(_file(path))
    assert result.parts[0].text.splitlines() == ["Heading", "A\tB", "Body"]


def test_excel_keeps_zero_and_formula_text(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append([0, "=1+1"])
    book.save(path)
    result = parse_local(_file(path))
    text = result.parts[0].text
    assert "0\t=1+1" in text


def test_powerpoint_reads_text_and_marks_image_only_slide(tmp_path: Path) -> None:
    path = tmp_path / "sample.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.text = "Hello"
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(str(path))
    result = parse_local(_file(path))
    assert result.parts[0].text == "Hello"
    assert result.parts[1].text == ""


def test_download_extension_uses_mime_for_dynamic_document_urls() -> None:
    assert extension_for("application/pdf", "https://example.test/download.aspx?id=7") == ".pdf"
    assert (
        extension_for(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "https://example.test/file?id=8",
        )
        == ".docx"
    )


PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def test_docx_visual_content_means_images(tmp_path: Path) -> None:
    image = tmp_path / "dot.png"
    image.write_bytes(PNG)

    def visual(name: str, *, text: bool, picture: bool) -> bool:
        doc = Document()
        if text:
            doc.add_paragraph("Nội dung")
        if picture:
            doc.add_picture(str(image))
        path = tmp_path / name
        doc.save(str(path))
        return parse_local(_file(path)).parts[0].has_visual_content

    assert visual("both.docx", text=True, picture=True) is True
    assert visual("text.docx", text=True, picture=False) is False
    assert visual("empty.docx", text=False, picture=False) is False


def test_pptx_visual_content_means_picture(tmp_path: Path) -> None:
    image = tmp_path / "dot.png"
    image.write_bytes(PNG)
    path = tmp_path / "deck.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Hello"
    slide.shapes.add_picture(str(image), Inches(1), Inches(3))
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(2), Inches(1)
    ).text = "Text only"
    deck.save(str(path))

    parts = parse_local(_file(path)).parts

    assert [part.has_visual_content for part in parts] == [True, False]


def test_rtf_goes_through_libreoffice(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "note.rtf"
    source.write_text(r"{\rtf1 Hello}", encoding="utf-8")
    targets: list[str] = []

    def convert(path: Path, out_dir: Path, target: str) -> Path:
        targets.append(target)
        converted = out_dir / "note.docx"
        doc = Document()
        doc.add_paragraph("Hello")
        doc.save(str(converted))
        return converted

    monkeypatch.setattr(dispatch, "convert_office", convert)

    result = parse_local(_file(source))

    assert targets == ["docx:Office Open XML Text"]
    assert result.parts[0].text == "Hello"
