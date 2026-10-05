from __future__ import annotations


import openpyxl
import xlrd

from ..models import DownloadedFile, ExtractedPart
from .common import result


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_xlsx(file: DownloadedFile):
    book_values = openpyxl.load_workbook(file.path, read_only=True, data_only=True)
    book_formulas = openpyxl.load_workbook(file.path, read_only=True, data_only=False)
    parts = []
    for index, name in enumerate(book_values.sheetnames, 1):
        values = book_values[name]
        formulas = book_formulas[name]
        rows = []
        for row_values, row_formulas in zip(values.iter_rows(), formulas.iter_rows()):
            cells = []
            for value, formula in zip(row_values, row_formulas):
                cached = _cell(value.value)
                cells.append(cached if cached else _cell(formula.value))
            line = "\t".join(cells).rstrip("\t")
            if line:
                rows.append(line)
        text = f"[{name}]\n" + "\n".join(rows) if rows else ""
        parts.append(ExtractedPart(index, "sheet", text, False, "openpyxl"))
    book_values.close()
    book_formulas.close()
    return result(file, parts, "openpyxl")


def parse_xls(file: DownloadedFile):
    book = xlrd.open_workbook(file.path, on_demand=True)
    parts = []
    for index, name in enumerate(book.sheet_names(), 1):
        sheet = book.sheet_by_name(name)
        rows = []
        for row in sheet.get_rows():
            line = "\t".join(_cell(cell.value) for cell in row).rstrip("\t")
            if line:
                rows.append(line)
        parts.append(
            ExtractedPart(
                index, "sheet", f"[{name}]\n" + "\n".join(rows) if rows else "", False, "xlrd"
            )
        )
    book.release_resources()
    return result(file, parts, "xlrd")
