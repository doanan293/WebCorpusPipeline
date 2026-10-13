from __future__ import annotations

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ir_data.models import DownloadedFile, ExtractedPart
from ir_data.parsers.common import result


def _shape_text(shape) -> list[str]:
    out = []
    if getattr(shape, "has_text_frame", False) and shape.text.strip():
        out.append(shape.text.strip())
    if getattr(shape, "has_table", False):
        out.append(
            "\n".join(
                "\t".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows
            )
        )
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            out.extend(_shape_text(child))
    return out


def parse_powerpoint(file: DownloadedFile):
    deck = Presentation(file.path)
    parts = []
    for index, slide in enumerate(deck.slides, 1):
        chunks = []
        image = False
        for shape in slide.shapes:
            chunks.extend(_shape_text(shape))
            image = image or shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        parts.append(
            ExtractedPart(index, "slide", "\n".join(chunks), image and not chunks, "python-pptx")
        )
    return result(file, parts, "python-pptx")
